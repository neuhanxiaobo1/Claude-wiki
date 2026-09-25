"""Authoritative workbook IO. No embedded literature facts or network access."""
from pathlib import Path
import json,hashlib,os
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils.cell import range_boundaries

HERE=Path(__file__).resolve().parent
BOOK=HERE/'BC_review_evidence.xlsx'
def digest(path=BOOK):return hashlib.sha256(path.read_bytes()).hexdigest()
def rows(ws):
    headers=[c.value for c in ws[1]]
    return [dict(zip(headers,values)) for values in ws.iter_rows(min_row=2,values_only=True) if any(v is not None for v in values)]
def validate(wb):
    schema=json.loads((HERE/'schema.json').read_text(encoding='utf-8'));data={};errors=[]
    ids={'Paper_Index':'Paper_ID','Evidence_Records':'Evidence_ID','Synthesis_Map':'Synthesis_ID'}
    for name,spec in schema['tables'].items():
        ws=wb[name];head=[c.value for c in ws[1]]
        if head!=spec['columns']:errors.append(name+': schema mismatch')
        rr=rows(ws);data[name]=rr;keys=[r[ids[name]] for r in rr]
        table=next(iter(ws.tables.values()))
        if len(rr)+1>range_boundaries(table.ref)[3]:errors.append(name+': data exceeds table range')
        if any(not k for k in keys) or len(keys)!=len(set(keys)):errors.append(name+': missing/duplicate ID')
        for dv in ws.data_validations.dataValidation:
            if dv.type!='list' or not dv.formula1.startswith('"'):continue
            allowed=dv.formula1.strip('"').split(',')
            for bounds in dv.sqref.ranges:
                for row in ws.iter_rows(min_row=2,max_row=len(rr)+1,min_col=bounds.min_col,max_col=bounds.max_col):
                    for c in row:
                        if c.value is not None and c.value not in allowed:errors.append(f'{name}!{c.coordinate}: invalid enum')
    papers={r['Paper_ID']:r for r in data['Paper_Index']};ev={r['Evidence_ID']:r for r in data['Evidence_Records']}
    required=['Paper_ID','Track','Review_Question','Sample_or_Group','Intervention_or_Condition','Source_Location','Review_Usable_Claim','Evidence_Type','Verification_Status','Verification_Scope','Support_Assessment','Limitation','Data_Group_ID']
    for r in ev.values():
        for field in required:
            if not r.get(field):errors.append(f'{r["Evidence_ID"]}: missing {field}')
        if r['Paper_ID'] not in papers:errors.append(r['Evidence_ID']+': unknown paper')
        if r.get('Needs_Check')=='Yes' and not r.get('Check_Reason'):errors.append(r['Evidence_ID']+': check reason missing')
        if r.get('Evidence_Type')=='Author-interpretation' and r.get('Evidence_Directness')!='Author-interpretation':errors.append(r['Evidence_ID']+': interpretation mislabeled')
        for ref in splitids(r.get('Related_Evidence_IDs')):
            if ref not in ev:errors.append(r['Evidence_ID']+': unknown related evidence '+ref)
    for r in data['Synthesis_Map']:
        supports=splitids(r.get('Supporting_Evidence_IDs'));limits=splitids(r.get('Limiting_Evidence_IDs'))
        if not supports:errors.append(r['Synthesis_ID']+': evidence IDs missing')
        for ref in supports+limits:
            if ref not in ev:errors.append(r['Synthesis_ID']+': unknown evidence '+ref)
        cited=set(splitids(r.get('Supporting_Paper_IDs'))+splitids(r.get('Contradicting_or_Limiting_Paper_IDs')))
        actual={ev[e]['Paper_ID'] for e in supports+limits if e in ev}
        if cited!=actual:errors.append(r['Synthesis_ID']+': paper/evidence mapping mismatch')
        if len(actual)<2:errors.append(r['Synthesis_ID']+': fewer than two papers')
        if not r.get('Strength_Rationale') or not r.get('Comparability'):errors.append(r['Synthesis_ID']+': evaluation missing')
        if r.get('Readiness')=='Ready' and any(ev[e]['Verification_Status']!='Checked' for e in supports if e in ev):errors.append(r['Synthesis_ID']+': unchecked Ready basis')
    return errors,data
def splitids(value):return [s.strip() for s in str(value or '').split(';') if s.strip() and s.strip()!='NA']
def append_batch(updates,expected_hash):
    if digest()!=expected_hash:raise RuntimeError('Workbook drift; refusing overwrite')
    wb=openpyxl.load_workbook(BOOK)
    for name,records in updates.items():
        ws=wb[name];headers=[c.value for c in ws[1]];start=len(rows(ws))+2
        if start+len(records)-1>range_boundaries(next(iter(ws.tables.values())).ref)[3]:raise ValueError('Extend table and validations before adding these rows')
        for n,record in enumerate(records,start):
            if set(record)-set(headers):raise ValueError('Unknown columns '+str(set(record)-set(headers)))
            for col,header in enumerate(headers,1):
                cell=ws.cell(n,col,record.get(header));cell.alignment=Alignment(wrap_text=True,vertical='top')
            ws.row_dimensions[n].height=100
    errors,_=validate(wb)
    if errors:raise ValueError(errors)
    temp=BOOK.with_name('BC_review_evidence.pending.xlsx')
    if temp.exists():raise RuntimeError('Pending workbook exists; inspect it before proceeding')
    wb.save(temp)
    reread=openpyxl.load_workbook(temp);errors,data=validate(reread);reread.close()
    if errors:raise ValueError(errors)
    if digest()!=expected_hash:raise RuntimeError('Workbook changed during save; candidate retained')
    os.replace(temp,BOOK)
    return {k:len(v) for k,v in data.items()}
