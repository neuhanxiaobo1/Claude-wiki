"""Build an EMPTY candidate only; never edit source or an existing candidate.
Run: python -B directions/wave-transparent-composites/docs/bc-integration-2026-09-25/build_candidate.py
Requires the original source at its recorded path and openpyxl. No network.
"""
from pathlib import Path
from copy import copy
import hashlib
import json
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sources = json.loads((HERE / 'sources.json').read_text(encoding='utf-8'))
source = next(s for s in sources if s['name'].endswith('.xlsx'))
target = HERE / 'BC_review_evidence_v2_candidate.xlsx'
if target.exists():
    raise SystemExit('Candidate exists; refusing to overwrite.')
for s in sources:
    assert hashlib.sha256(Path(s['path']).read_bytes()).hexdigest() == s['sha256'], s['name']

# name, definition, allowed values (empty means free text)
extra = {
    'Paper_Index': [
        ('B_Use', 'B的论证用途；相关度不等于用途，Core不代表已核查。', ['Core','Context','Method','Review','Exclude']),
        ('C_Use', 'C的论证用途，与B独立判断。', ['Core','Context','Method','Review','Exclude']),
        ('Zotero_Item_Key', '条目键；有多文库时同时注明libraryID。', []),
        ('Legacy_Wiki_Page', '相对vault根的完整旧paper路径；无旧页写NA。', []),
        ('Source_Version', '实际所读主PDF/MD/SI的版本与来源manifest或哈希；不是文献出版年份。', []),
        ('Verification_Scope', 'Read_Status=Verified时必须说明实际核查范围及残余缺项。', []),
    ],
    'Evidence_Records': [
        ('Evidence_Type', '本行主要断言类型，观察与解释拆行并关联。', ['Observation','Model-output','Author-interpretation','Secondary-citation','AI-inference']),
        ('Verification_Status', '仅本条的核查状态，不替代paper review_status。', ['Checked','Partial','Pending','Unavailable']),
        ('Verification_Scope', '核过的原文材料/定位及未核范围，不能只写已读。', []),
        ('Check_Reason', 'Needs_Check=Yes时说明问题及影响；未解决来源矛盾保留各位置。', []),
        ('Pairing_Status', '针对本行关系：同试样、同批匹配、不同状态、未知或不适用。不能单独证明控制变量。', ['Same-specimen','Matched-batch','Different-state','Unclear','NA']),
        ('Data_Group_ID', '独立性分组标识；共用原始数据的多篇论文共用组。未知写Unknown并限制独立验证计数。', []),
        ('Related_Evidence_IDs', '关联EV编号，以分号分隔；配对关系及角色在本行条件/限制列说明。', []),
        ('Legacy_Evidence_Ref', '完整旧paper路径及实际E标题/块锚点；多项分号分隔；无旧证据写NA。', []),
        ('Support_Assessment', '相对于Review_Usable_Claim的支持程度；理由写Limitation（含支持理由）。', ['Sufficient','Partial','Insufficient']),
    ],
    'Synthesis_Map': [
        ('Supporting_Evidence_IDs', '精确支持本句的EV编号，分号分隔；须能回到Supporting_Paper_IDs。', []),
        ('Limiting_Evidence_IDs', '限制/反对本句的EV编号，分号分隔；不得省略已知实质反证。无记录不等于无反例。', []),
        ('Comparability', '本次综合可比性；理由与条件写Common_Conditions/Key_Differences。', ['Comparable','Conditional','Qualitative-only','Not-comparable']),
        ('Strength_Rationale', '证据强度相对于Candidate_Claim的理由，含独立性、未决影响及适用范围。', []),
        ('Readiness', 'Ready仅当前限定句子可用；Provisional内部分析；Blocked不进确定性正文。', ['Ready','Provisional','Blocked']),
    ],
}

wb = openpyxl.load_workbook(source['path'])
schema = {
    'direction_id': 'wave-transparent-composites', 'version': '2026-09-25-candidate',
    'status': 'empty-design-candidate-not-production', 'source_sha256': source['sha256'],
    'id_policy': {'paper': 'P0001, never reuse', 'evidence': 'EV000001, never reuse',
                  'synthesis': 'SYN-B-001 / SYN-C-001, never reuse',
                  'multi_id_separator': ';', 'legacy_evidence': 'retain original page-local E#'},
    'tables': {},
    'validation_limits': ['Excel dropdowns do not enforce pasted/programmatic values.',
                          'Required fields depend on row stage; empty template rows are allowed.',
                          'Foreign keys, DOI identity, source locations, sample pairing and scientific support require separate checks.'],
}
for name, fields in extra.items():
    ws = wb[name]
    old_headers = [c.value for c in ws[1]]
    assert not any(c.value is not None for row in ws.iter_rows(min_row=2) for c in row), name
    table = next(iter(ws.tables.values()))
    for col, (field, meaning, options) in enumerate(fields, len(old_headers)+1):
        cell = ws.cell(1, col, field)
        cell._style = copy(ws.cell(1, len(old_headers))._style)
        cell.comment = openpyxl.comments.Comment(meaning, 'BC design')
        ws.column_dimensions[get_column_letter(col)].width = 27 if options else 42
        for row in range(2, 1001):
            ws.cell(row, col)._style = copy(ws.cell(row, len(old_headers))._style)
        table.tableColumns.append(TableColumn(id=col, name=field))
        if options:
            dv = DataValidation(type='list', formula1='"'+','.join(options)+'"', allow_blank=True)
            dv.add(f'{get_column_letter(col)}2:{get_column_letter(col)}1000')
            ws.add_data_validation(dv)
    last = get_column_letter(ws.max_column)
    table.ref = f'A1:{last}1000'
    if table.autoFilter:
        table.autoFilter.ref = table.ref
    for dv in ws.data_validations.dataValidation:
        dv.sqref = str(dv.sqref).replace('2000', '1000')
        dv.allowBlank = True
        dv.showErrorMessage = True
        dv.errorStyle = 'stop'
        dv.errorTitle = '值不在候选范围'
        dv.error = '请选择列表值；粘贴或程序写入后仍需独立校验。'
    ws.freeze_panes = 'C2'
    schema['tables'][name] = {'original_headers': old_headers, 'columns': [c.value for c in ws[1]],
        'added_fields': [{'name':f,'meaning':m,'allowed_values':v} for f,m,v in fields],
        'table_ref':table.ref}

instructions = [
    ('项目', '说明'),
    ('状态', '2026-09-25空白候选；未导入文献，未正式启用。B/C并行验证，不沿用C优先。'),
    ('三主表', 'Paper_Index每论文一行；Evidence_Records每主要可引用断言一行；Synthesis_Map跨论文论证。'),
    ('唯一维护源', '正式接入后Excel维护BC事实，旧paper保留历史地址；原始PDF仍为科学依据。'),
    ('证据与Both', '同一事实只一条EV；Both供B/C引用，观测与机制解释拆行并关联。'),
    ('用途', 'B_Use/C_Use独立于相关度；Core不等于核查通过或机制成立。'),
    ('核查', '类型、核查状态、支持程度分别填写；Verified说明范围，不据局部Checked升级整篇。'),
    ('复核', '非关键疑点可集中复核；改变论断的身份/单位/状态冲突先核或隔离受影响判断。'),
    ('配对', '同论文不等于同状态；写清样品/热状态/测量状态，关系引用Related_Evidence_IDs。'),
    ('综合', '至少两论文共同回答问题才建综合行，另核数据独立性；单篇事实无需伪造第二验证。'),
    ('引用', '综合必须引用具体EV并与Paper_ID对应；两轨总论文数去重Both。'),
    ('写作', '从综合组织论点、证据表提取事实；Ready仅当前限定表述可用，必要原文复核始终允许。'),
    ('工作量', '不额外写长篇单论文摘要；例外为复杂冲突/SI；不在阅读阶段写完整综述。'),
    ('缺失', '未报告写NR，不适用写NA，未知独立性写Unknown；不能用0代替缺失测量。'),
    ('旧ID', 'P0001 / EV000001 / SYN-B-001；旧paper E#保留，撤回记录留ID及说明，不复用。'),
    ('范围', '表和下拉均至1000行；超过时同步扩表和验证。空行不算文献。'),
    ('校验边界', '下拉不能保证粘贴数据正确；引用、必填条件、同源数据、科学证据另核。'),
    ('备份', '本候选xlsx当前被Git忽略；正式启用前设置本方向精确例外并核验实际跟踪。'),
    ('详细依据', '同目录report.md、schema.json及sources.json；本表为空，不作已入库证明。'),
]
ws = wb['Instructions']
ws.delete_rows(1, ws.max_row)
for row in instructions:
    ws.append(row)
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 105
for row in ws:
    for cell in row:
        cell.alignment = openpyxl.styles.Alignment(vertical='top', wrap_text=True)
    ws.row_dimensions[row[0].row].height = 43
wb.save(target)
(HERE/'schema.json').write_text(json.dumps(schema,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

check = openpyxl.load_workbook(target)
results = {}
for name, spec in schema['tables'].items():
    ws = check[name]
    assert [c.value for c in ws[1]] == spec['columns']
    assert len(set(spec['columns'])) == len(spec['columns'])
    count = sum(any(c.value is not None for c in row) for row in ws.iter_rows(min_row=2))
    assert count == 0
    table = next(iter(ws.tables.values()))
    assert table.ref == spec['table_ref']
    assert [c.name for c in table.tableColumns] == spec['columns']
    for dv in ws.data_validations.dataValidation:
        assert dv.showErrorMessage and dv.allowBlank
        assert all(r.max_row == 1000 for r in dv.sqref.ranges)
    results[name] = {'data_rows':count,'columns':len(spec['columns']),'table_ref':table.ref,
                     'validation_count':len(ws.data_validations.dataValidation)}
for s in sources:
    assert hashlib.sha256(Path(s['path']).read_bytes()).hexdigest() == s['sha256']
results.update({'external_sources_unchanged':True, 'original_headers_preserved':True,
                'candidate_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
                'scientific_evidence_checked':False, 'production_activated':False})
(HERE/'candidate-validation.json').write_text(json.dumps(results,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(results,ensure_ascii=False))
