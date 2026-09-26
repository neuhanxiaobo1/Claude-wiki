"""In-memory negative checks for evidence integrity. Never saves the workbook."""
import json,io
import openpyxl
from workbook_io import BOOK,rows,validate,digest

def setcell(w,sheet,row,column,value):
    ws=w[sheet];headers=[c.value for c in ws[1]];ws.cell(row,headers.index(column)+1,value)

def run():
    before=digest();source=BOOK.read_bytes();results=[]
    def check(name,change,expected):
        w=openpyxl.load_workbook(io.BytesIO(source));change(w);errors,_=validate(w);w.close()
        ok=any(expected in e for e in errors)
        results.append({'case':name,'rejected':ok})
        if not ok:raise AssertionError((name,errors))
    check('duplicate EV',lambda w:setcell(w,'Evidence_Records',3,'Evidence_ID','EV000001'),'duplicate ID')
    check('unknown paper',lambda w:setcell(w,'Evidence_Records',2,'Paper_ID','P9999'),'unknown paper')
    check('missing source manifest',lambda w:setcell(w,'Paper_Index',2,'Source_Manifest',''),'invalid source manifest')
    check('out-of-direction source',lambda w:setcell(w,'Paper_Index',2,'Source_Manifest','AGENTS.md'),'invalid source manifest')
    check('unchecked Ready',lambda w:setcell(w,'Synthesis_Map',2,'Readiness','Ready') or setcell(w,'Evidence_Records',2,'Verification_Status','Partial'),'unchecked Ready')
    check('unclassified cannot be Checked',lambda w:setcell(w,'Evidence_Records',2,'Evidence_Type','Unclassified-report'),'unclassified report')
    check('interpretation is not observation',lambda w:setcell(w,'Evidence_Records',2,'Evidence_Type','Author-interpretation'),'interpretation mislabeled')
    check('unknown related EV',lambda w:setcell(w,'Evidence_Records',2,'Related_Evidence_IDs','EV999999'),'unknown related evidence')
    if digest()!=before:raise AssertionError('Workbook was changed by validation')
    return {'source_sha256':before,'negative_checks':results,'workbook_unchanged':True,'scientific_correctness_tested':False}
if __name__=='__main__':print(json.dumps(run(),indent=2))
